import io
import magic
from typing import Tuple, Optional

import pdfplumber
from docx import Document
import pypdf2

from backend.utils.file_utils import (
    FileParsingError,
    TextExtractionError,
    FileUploadError,
    log_error,
    log_warning,
    log_info,
    with_fallback
)

from backend.core.config import (
    MAX_FILE_SIZE_BYTES,
    MAX_FILE_SIZE_MB,
    SUPPORTED_MIME_TYPES,
)

class FileParsingError:
    pass

class FileValidationError:
    pass

def validate_file(file_data: bytes, filename: str) -> Tuple[bool, str, Optional[str]]:
    file_size_bytes = len(file_data)
    if file_size_bytes > MAX_FILE_SIZE_BYTES:
        size_mb = file_size_bytes / (1024 * 1024)
        return False, (
            f'File size ({size_mb:.2f} MB) exceeds the maximum of {MAX_FILE_SIZE_MB} MB.',
            'Please upload a smaller file or compress your resume.',
        ), None

    if file_size_bytes == 0:
        return False, (
            'The uploaded file is empty.',
            'Please upload a valid file.',
        ), None

    try:
        mime_type = magic.from_buffer(file_data, mime = True)
    except Exception as e:
        return False, f"Error determining file type: {e}", None

    if mime_type not in SUPPORTED_MIME_TYPES:
        supported_types = ", ".join(SUPPORTED_MIME_TYPES.keys())

        return False, (
            f'Unsupported file type {mime_type}. Allowed types: {supported_types}',
            'Please upload a valid file type.',
        ), None

    return True, '', mime_type

def _extract_pdf_hyperlinks(file_data: bytes) -> str:
    urls = []
    try:
        reader = pypdf2.PdfReader(io.BytesIO(file_data))
        for page in reader.pages:
            if '/Annots' not in page:
                continue
            for annot_ref in page['/Annots']:
                try:
                    annot = annot_ref.get_object()
                    if annot.get('/Subtype') != "/Link":
                        continue
                    action = annot.get('/A', {})
                    uri = action.get('/URI', '')
                    if uri and isinstance(uri, (str, bytes)):
                        # PyPDF2 may return bytes for URI values
                        if isinstance(uri, bytes):
                            uri = uri.decode('utf-8', errors='ignore')
                        uri = uri.strip()
                        if uri.startswith('http'):
                            urls.append(uri)
                except Exception:
                    pass
    except Exception:
        log_warning(f"Error extracting hyperlinks from PDF: {e}")
    return "\n".join(urls)

def _extract_pdf_with_pdfplumber(file_data: bytes) -> str:
    text = ''
    try:
        with pdfplumber.open(io.BytesIO(file_data)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        
        if not text.strip():
            raise TextExtractionError(
                'pdfplumber failed to extract text from PDF',
                user_message = 'No text could be extracted from the PDF file.',
            )

        hyperlinks = _extract_pdf_hyperlinks(file_data)
        if hyperlinks:
            text += f"\n\nHyperlinks:\n{hyperlinks}"

        return text.strip()
    except Exception as e:
        raise TextExtractionError(
            f"Error extracting text from PDF: {e}",
            user_message = 'An error occurred while extracting text from the PDF file.',
        )
    
def _extract_pdf_with_pypdf2(file_data: bytes) -> str:
    text = ''
    try:
        pdf_reader = pypdf2.PdfReader(io.BytesIO(file_data))
        for page in pdf_reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        
        if not text.strip():
            raise TextExtractionError(
                'pypdf2 failed to extract text from PDF',
                user_message = 'No text could be extracted from the PDF file.',
            )

        hyperlinks = _extract_pdf_hyperlinks(file_data)
        if hyperlinks:
            text += f"\n\nHyperlinks:\n{hyperlinks}"

        return text.strip()
    except Exception as e:
        raise TextExtractionError(
            f"Error extracting text from PDF: {e}",
            user_message = 'An error occurred while extracting text from the PDF file.',
        )

def extract_text_from_pdf(file_data: bytes) -> str:
    try:
        result, used_fallback = with_fallback(
            _extract_pdf_with_pdfplumber,
            _extract_pdf_with_pypdf2,
            file_data,
            log_fallback = True
        )
        if used_fallback:
            log_info(f"Used fallback method for PDF extraction: {used_fallback.__name__}", context ='resume_parser')
        return result
    except Exception as e:
        log_error(f"Error extracting text from PDF: {e}", context ='extract_text_from_pdf')
        raise TextExtractionError(
            f"Error extracting text from PDF: {e}",
            user_message = 'An error occurred while extracting text from the PDF file.',
        )


def extract_text_from_docx(file_data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_data))
        text_parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        text = '\n'.join(text_parts)

        if not text.strip():
            raise FileParsingError(
                'No text could be extracted from the document. '
                'The document may be empty or corrupted.'
            )
        
        try:
            for rel in doc.part.rels.values():
                if 'hyperlink' in rel.reltype.lower():
                    url = rel._target
                    if isinstance(url, str) and url.startswith('http'):
                        text += '\n' + url
        except Exception:
            pass

        log_info(f'Extracted {len(text)} chars from DOCX', context='resume_parser')
        return text.strip()

    except FileParsingError:
        raise   # Re-raise unchanged — don't wrap in another FileParsingError

    except Exception as e:
        log_error(e, context='extract_text_from_docx')
        raise FileParsingError(
            'Failed to extract text from DOCX. '
            'The document may be corrupted or in an unsupported format. '
            'Please try re-saving or converting to PDF.'
        ) from e

def extract_text_from_doc(file_data: bytes) -> str:
    raise FileParsingError(
        'Legacy .doc format is not supported. '
        'Please convert your document to .docx or .pdf and try again. '
        'You can convert using Microsoft Word, Google Docs, or online tools.'
    )

def extract_text(file_data:bytes, file_type:str)->str:
    if file_type=='pdf':
        return extract_text_from_pdf(file_data)
    elif file_type=='docx':
        return extract_text_from_docx(file_data)
    elif file_type=='doc':
        return extract_text_from_doc(file_data)
    else:
        raise FileValidationError(
            f'invalid file type: {file_type}. supported types are: pdf, docx and doc',
            user_message = 'Please upload a valid file type.',
        )

def parse_resume(file_data:bytes, file_type:str)->Tuple[str, Dict]:
    log_info(f"Parsing resume: {file_type}", context ='parse_resume')
    
    #phase01: validate file
    try:
        is_valid, error_message, mime_type = validate_file(file_data, file_type)
        if not is_valid:
            log_error(error_message, context ='parse_resume')
            raise FileValidationError(
                error_message,
                user_message = 'Please upload a valid file.',
            )
    except FileValidationError:
        raise   # Re-raise unchanged — don't wrap in another FileValidationError

    except Exception as e:
        log_error(f"Error validating file: {e}", context ='parse_resume')
        raise FileValidationError(
            f"Error validating file: {e}",
            user_message = 'Please upload a valid file.',
        ) from e

    #phase02: extraction of file
    try:
        text = extract_text(file_data, file_type)
        log_info(f'Extracted {len(text)} chars from {filename}', context='parse_resume_file')

    except FileParsingError:
        raise   # Re-raise unchanged

    except Exception as e:
        log_error(e, context='parse_resume_file_extraction')
        raise FileParsingError(
            'An unexpected error occurred while processing the file. '
            'Please try again or contact support if the problem persists.'
        ) from e

    metadata = {
        'filename':        filename,
        'file_type':       file_type,
        'file_size_bytes': len(file_data),
        'text_length':     len(text),
        'success':         True,
    }
    return text, metadata
